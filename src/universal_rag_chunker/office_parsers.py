"""Optional binary-format adapters with explicit extraction scope and locators."""
from __future__ import annotations

import hashlib
from importlib import import_module, metadata
from io import BytesIO
from datetime import date, datetime, time
from pathlib import Path
import re

from .models import Block, DocumentIR, VERSION, digest


def dependency(module, distribution, extra):
    try:
        return import_module(module), metadata.version(distribution)
    except ImportError as error:
        raise ValueError(f"MISSING_DEPENDENCY: install with python -m pip install 'universal-rag-chunker[{extra}]'") from error


class BinaryParser:
    version = VERSION

    def create(self, source: Path, source_id: str, raw: bytes, library_version: str, options: dict):
        parser = {"name": self.name, "version": self.version,
                  "config_digest": digest({"library_version": library_version, **options})}
        origin = {"source_id": source_id, "display_name": source.name, "media_type": self.media_type,
                  "content_sha256": hashlib.sha256(raw).hexdigest()}
        document_id = "doc_" + digest([source_id, origin["content_sha256"], parser])
        return DocumentIR(document_id, origin, parser, [], metadata={"extraction": options, "diagnostics": []})

    @staticmethod
    def append(document, kind, text, locator, attributes=None, parent_id=None):
        attributes = attributes or {}
        order = len(document.blocks)
        block_id = "block_" + digest([document.document_id, order, kind, text, locator, parent_id, attributes])
        block = Block(block_id, kind, order, parent_id, text, locator, attributes)
        document.blocks.append(block)
        return block

    @staticmethod
    def finish(document):
        document.source["extracted_text_sha256"] = hashlib.sha256(
            "".join(b.text for b in document.blocks).encode("utf-8")).hexdigest()
        if not any(b.text.strip() for b in document.blocks):
            raise ValueError("NO_EXTRACTABLE_TEXT: file contains no text in supported extraction scope")
        return document


class PdfParser(BinaryParser):
    name = "pdf"
    suffixes = (".pdf",)
    media_type = "application/pdf"

    def parse(self, source: Path, source_id: str):
        pdf, version = dependency("pypdf", "pypdf", "pdf")
        raw = source.read_bytes()
        document = self.create(source, source_id, raw, version, {"scope": "page_text", "ocr": False})
        try:
            reader = pdf.PdfReader(BytesIO(raw))
            if reader.is_encrypted:
                raise ValueError("ENCRYPTED_PDF: provide an unencrypted copy")
            document.metadata["page_count"] = len(reader.pages)
            for number, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                self.append(document, "paragraph", text, {"kind": "page", "page": number})
                if not text.strip():
                    document.metadata["diagnostics"].append({"code": "PDF_PAGE_NO_TEXT", "severity": "warning",
                        "page": number, "message": "Page has no extractable text; may be blank or require OCR."})
            if not any(b.text.strip() for b in document.blocks):
                raise ValueError("PDF_NO_TEXT: no text layer found; scanned PDFs require OCR (not included)")
        except Exception as error:
            raise ValueError(f"PDF extraction failed: {error}") from error
        return self.finish(document)


class DocxParser(BinaryParser):
    name = "docx"
    suffixes = (".docx",)
    media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    def parse(self, source: Path, source_id: str):
        docx, version = dependency("docx", "python-docx", "word")
        from docx.text.paragraph import Paragraph
        raw = source.read_bytes()
        document = self.create(source, source_id, raw, version, {"scope": "body_paragraphs_and_tables"})
        hierarchy = []
        try:
            word = docx.Document(BytesIO(raw))
            for body_index, item in enumerate(word.iter_inner_content(), 1):
                locator = {"kind": "docx", "body_index": body_index}
                if isinstance(item, Paragraph):
                    text = item.text + "\n"
                    style = item.style.style_id if item.style else ""
                    heading = re.fullmatch(r"Heading([1-6])", style or "")
                    attrs = {"style": style}
                    kind = "paragraph"
                    if heading and item.text.strip():
                        kind = "heading"
                        level = int(heading[1])
                        attrs.update(level=level, title=item.text)
                        while hierarchy and hierarchy[-1][0] >= level:
                            hierarchy.pop()
                    block = self.append(document, kind, text, locator, attrs, hierarchy[-1][1] if hierarchy else None)
                    if kind == "heading":
                        hierarchy.append((level, block.block_id))
                else:
                    seen = set()
                    for row_number, row in enumerate(item.rows, 1):
                        cells = []
                        for col_number, cell in enumerate(row.cells, 1):
                            # python-docx exposes a merged cell repeatedly in the grid.
                            if cell._tc in seen:
                                continue
                            seen.add(cell._tc)
                            value = cell_text(cell)
                            if value:
                                cells.append({"column": col_number, "text": value})
                        if cells:
                            text = "".join(f"R{row_number}C{c['column']}: {c['text']}\n" for c in cells)
                            self.append(document, "table_row", text, {**locator, "row": row_number},
                                        {"cells": cells}, hierarchy[-1][1] if hierarchy else None)
            document.metadata["diagnostics"].append({"code": "DOCX_BODY_SCOPE", "severity": "info",
                "message": "Extracted body paragraphs/tables. Images, text boxes, headers, footers and tracked revisions are outside scope."})
        except Exception as error:
            raise ValueError(f"DOCX extraction failed: {error}") from error
        return self.finish(document)


def cell_text(cell):
    """Keep nested table text in its parent cell, in body order."""
    from docx.text.paragraph import Paragraph
    parts = []
    for item in cell.iter_inner_content():
        if isinstance(item, Paragraph):
            if item.text:
                parts.append(item.text)
        else:
            seen = set()
            for row in item.rows:
                for nested in row.cells:
                    if nested._tc not in seen:
                        seen.add(nested._tc)
                        value = cell_text(nested)
                        if value:
                            parts.append(value)
    return "\n".join(parts)


def render_value(value):
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


class XlsxParser(BinaryParser):
    name = "xlsx"
    suffixes = (".xlsx",)
    media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def parse(self, source: Path, source_id: str):
        excel, version = dependency("openpyxl", "openpyxl", "excel")
        raw = source.read_bytes()
        document = self.create(source, source_id, raw, version,
                               {"scope": "all_worksheets_nonempty_cells", "formula_mode": "formula_text"})
        workbook = None
        try:
            workbook = excel.load_workbook(BytesIO(raw), read_only=False, data_only=False)
            sheets = []
            for sheet in workbook.worksheets:
                sheets.append({"name": sheet.title, "state": sheet.sheet_state,
                               "merged_ranges": [str(r) for r in sheet.merged_cells.ranges]})
                for row in sheet.iter_rows():
                    cells = []
                    for cell in row:
                        if cell.value is None or cell.value == "":
                            continue
                        cells.append({"coordinate": cell.coordinate, "column": cell.column,
                                      "value": render_value(cell.value), "data_type": cell.data_type})
                    if not cells:
                        continue
                    text = "".join(f"{cell['coordinate']}: {cell['value']}\n" for cell in cells)
                    self.append(document, "table_row", text,
                                {"kind": "sheet", "sheet_name": sheet.title, "row": row[0].row,
                                 "column_start": min(c["column"] for c in cells),
                                 "column_end": max(c["column"] for c in cells)}, {"cells": cells})
            document.metadata["worksheets"] = sheets
            document.metadata["diagnostics"].append({"code": "XLSX_FORMULAS_PRESERVED", "severity": "info",
                "message": "Formulas are preserved as text, not calculated. All sheets including hidden sheets are included; merged-cell values occur at the anchor only."})
        except Exception as error:
            raise ValueError(f"XLSX extraction failed: {error}") from error
        finally:
            if workbook is not None:
                workbook.close()
        return self.finish(document)
