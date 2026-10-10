"""Real binary-format fixtures; no external documents or network required."""
import contextlib
import copy
import importlib.util
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from universal_rag_chunker import Config, Pipeline
from universal_rag_chunker.cli import main
from universal_rag_chunker.exporters import export
from universal_rag_chunker.pipeline import ParseError
from universal_rag_chunker.validators import ValidationError, validate_ir


def write_pdf(path, texts, encrypted=False):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    font = DictionaryObject({NameObject('/Type'):NameObject('/Font'), NameObject('/Subtype'):NameObject('/Type1'),
                             NameObject('/BaseFont'):NameObject('/Helvetica')})
    for text in texts:
        page = writer.add_blank_page(width=612,height=792)
        if text:
            page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
            stream = DecodedStreamObject()
            escaped=text.replace('\\','\\\\').replace('(','\\(').replace(')','\\)')
            stream.set_data(('BT /F1 12 Tf 50 700 Td ('+escaped+') Tj ET').encode('ascii'))
            page[NameObject('/Contents')] = stream
    if encrypted:
        writer.encrypt('test-password')
    writer.write(path)
    writer.close()


class OfficeTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name)

    def run_file(self,path,**options):
        return Pipeline().run(path,Config(**options))

    @unittest.skipUnless(importlib.util.find_spec('pypdf'),'Install [pdf] extra')
    def test_pdf_pages_and_determinism(self):
        source=self.root/'pages.pdf'
        write_pdf(source,['First page text.','Second page text.'])
        first=self.run_file(source)
        second=self.run_file(source)
        self.assertEqual(first.document.to_dict(),second.document.to_dict())
        self.assertEqual([c.to_dict() for c in first.chunks],[c.to_dict() for c in second.chunks])
        self.assertEqual([c.metadata['core.source_locations'][0]['page'] for c in first.chunks],[1,2])
        self.assertEqual(first.report['coverage_scope'],'parser_extracted_text')
        self.assertEqual(first.report['coverage'],1)

    @unittest.skipUnless(importlib.util.find_spec('pypdf'),'Install [pdf] extra')
    def test_pdf_blank_page_warning_and_page_identity(self):
        source=self.root/'mixed.pdf'
        write_pdf(source,['', 'Visible text'])
        result=self.run_file(source)
        self.assertEqual(len(result.document.blocks),2)
        self.assertEqual(result.chunks[0].metadata['core.source_locations'][0]['page'],2)
        self.assertEqual(result.report['diagnostics'][0]['code'],'PDF_PAGE_NO_TEXT')

    @unittest.skipUnless(importlib.util.find_spec('pypdf'),'Install [pdf] extra')
    def test_pdf_no_text_and_encryption(self):
        for encrypted, code in ((False,'PDF_NO_TEXT'),(True,'ENCRYPTED_PDF')):
            source=self.root/'problem.pdf'
            write_pdf(source,[''],encrypted=encrypted)
            with self.assertRaisesRegex(ParseError,code):
                self.run_file(source)

    @unittest.skipUnless(importlib.util.find_spec('docx'),'Install [word] extra')
    def test_docx_body_table_order_and_headings(self):
        from docx import Document
        word=Document()
        word.add_heading('使用指南',1)
        word.add_paragraph('前面的段落。')
        table=word.add_table(rows=2,cols=2)
        table.cell(0,0).text='名称'
        table.cell(0,1).text='值'
        table.cell(1,0).text='样例'
        table.cell(1,1).text='123'
        word.add_paragraph('后面的段落。')
        source=self.root/'guide.docx'
        word.save(source)
        result=self.run_file(source)
        blocks=result.document.blocks
        self.assertEqual([b.type for b in blocks],['heading','paragraph','table_row','table_row','paragraph'])
        self.assertEqual(blocks[2].locator,{'kind':'docx','body_index':3,'row':1})
        self.assertEqual(blocks[-1].parent_id,blocks[0].block_id)
        self.assertIn('名称',blocks[2].text)
        self.assertEqual(result.report['coverage'],1)

    @unittest.skipUnless(importlib.util.find_spec('docx'),'Install [word] extra')
    def test_docx_merged_cells_and_nested_table(self):
        from docx import Document
        word=Document()
        table=word.add_table(rows=2,cols=2)
        table.cell(0,0).merge(table.cell(0,1)).text='合并单元格'
        table.cell(1,0).text='普通'
        nested=table.cell(1,1).add_table(rows=1,cols=1)
        nested.cell(0,0).text='嵌套内容'
        source=self.root/'merged.docx'
        word.save(source)
        result=self.run_file(source)
        text=''.join(b.text for b in result.document.blocks)
        self.assertEqual(text.count('合并单元格'),1)
        self.assertIn('嵌套内容',text)

    @unittest.skipUnless(importlib.util.find_spec('openpyxl'),'Install [excel] extra')
    def test_xlsx_all_sheets_two_columns_and_types(self):
        from openpyxl import Workbook
        from datetime import datetime
        book=Workbook()
        sheet=book.active
        sheet.title='登记'
        sheet.append(['名称','值'])
        sheet.append(['零',0])
        sheet.append(['真假',False])
        sheet.append(['日期',datetime(2026,10,10,12,0)])
        sheet.append(['公式','=SUM(B2:B2)'])
        sheet['A8']='空行之后'
        safety=book.create_sheet('事故记录')
        safety.append(['事故','保留'])
        hidden=book.create_sheet('隐藏表')
        hidden.sheet_state='hidden'
        hidden.append(['隐藏数据'])
        book.create_sheet('空白')
        source=self.root/'book.xlsx'
        book.save(source)
        book.close()
        result=self.run_file(source)
        self.assertEqual({b.locator['sheet_name'] for b in result.document.blocks},{'登记','事故记录','隐藏表'})
        text=''.join(b.text for b in result.document.blocks)
        for expected in ('B2: 0','B3: FALSE','2026-10-10T12:00:00','=SUM(B2:B2)','A8: 空行之后'):
            self.assertIn(expected,text)
        self.assertTrue(all(c.kind=='table' for c in result.chunks))
        self.assertEqual(result.report['coverage'],1)

    @unittest.skipUnless(importlib.util.find_spec('openpyxl'),'Install [excel] extra')
    def test_xlsx_merged_cells_single_column_and_long_cell(self):
        from openpyxl import Workbook
        book=Workbook()
        sheet=book.active
        sheet.merge_cells('A1:C1')
        sheet['A1']='合并标题'
        sheet['A2']='甲'*200
        source=self.root/'merged.xlsx'
        book.save(source)
        book.close()
        result=self.run_file(source,max_characters=30,minimum_characters=5,overlap_characters=3)
        self.assertEqual(result.document.blocks[0].text.count('合并标题'),1)
        self.assertTrue(all(len(c.text)<=30 for c in result.chunks))
        self.assertTrue(all(c.metadata['core.source_locations'][0]['row']==2 for c in result.chunks[1:]))

    def test_missing_optional_dependency_has_actionable_error(self):
        source=self.root/'source.pdf'
        source.write_bytes(b'pdf')
        with patch('universal_rag_chunker.office_parsers.import_module',side_effect=ImportError('not installed')):
            with self.assertRaisesRegex(ParseError,'MISSING_DEPENDENCY'):
                self.run_file(source)

    def test_legacy_formats_are_not_claimed(self):
        for suffix in ('.doc','.xls'):
            source=self.root/('old'+suffix)
            source.write_bytes(b'unsupported')
            with self.assertRaises(ValueError):
                self.run_file(source)

    def test_corrupt_files_fail_without_traceback(self):
        for suffix,module in (('.pdf','pypdf'),('.docx','docx'),('.xlsx','openpyxl')):
            if not importlib.util.find_spec(module):
                continue
            source=self.root/('broken'+suffix)
            source.write_bytes(b'not a real binary file')
            with contextlib.redirect_stdout(io.StringIO()),contextlib.redirect_stderr(io.StringIO()):
                self.assertEqual(main(['chunk',str(source),'--output',str(self.root/'out')]),3)
            self.assertFalse((self.root/'out').exists())

    @unittest.skipUnless(importlib.util.find_spec('pypdf'),'Install [pdf] extra')
    def test_office_export_validate_and_schema_tampering(self):
        source=self.root/'page.pdf'
        write_pdf(source,['A very long text. '*20])
        result=self.run_file(source,max_characters=40,minimum_characters=10,overlap_characters=5)
        output=self.root/'out'
        export(result,output)
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(main(['validate',str(output/'document.ir.json'),'--chunks',str(output/'chunks.jsonl'),
                                   '--max-characters','40']),0)
        bad=copy.deepcopy(result.document)
        bad.blocks[0].locator['page']=0
        with self.assertRaises(ValidationError):
            validate_ir(bad)
        bad=copy.deepcopy(result.document)
        bad.source['extracted_text_sha256']='0'*64
        with self.assertRaises(ValidationError):
            validate_ir(bad)

    def test_old_schema_still_validates(self):
        source=self.root/'legacy.txt'
        source.write_bytes(b'Legacy text')
        result=self.run_file(source)
        result.document.schema_version='1.0'
        for chunk in result.chunks:
            chunk.schema_version='1.0'
        output=self.root/'legacy-out'
        export(result,output)
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(main(['validate',str(output/'document.ir.json'),'--chunks',str(output/'chunks.jsonl')]),0)

    @unittest.skipUnless(importlib.util.find_spec('docx') and importlib.util.find_spec('openpyxl'),'Install [word,excel] extras')
    def test_word_excel_determinism_and_cli_round_trip(self):
        from docx import Document
        from openpyxl import Workbook
        word=Document()
        word.add_heading('标题',1)
        word.add_paragraph('真实文档内容。')
        docx=self.root/'real.docx'
        word.save(docx)
        book=Workbook()
        book.active.append(['名称','内容'])
        book.active.append(['示例','实际值'])
        xlsx=self.root/'real.xlsx'
        book.save(xlsx)
        book.close()
        for source in (docx,xlsx):
            first=self.run_file(source)
            second=self.run_file(source)
            self.assertEqual(first.document.to_dict(),second.document.to_dict())
            self.assertEqual([c.to_dict() for c in first.chunks],[c.to_dict() for c in second.chunks])
            output=self.root/source.suffix[1:]
            with contextlib.redirect_stdout(io.StringIO()),contextlib.redirect_stderr(io.StringIO()):
                self.assertEqual(main(['chunk',str(source),'--output',str(output),'--format','json']),0)
                self.assertEqual(main(['validate',str(output/'document.ir.json'),'--chunks',str(output/'chunks.json')]),0)


if __name__=='__main__':
    unittest.main()
